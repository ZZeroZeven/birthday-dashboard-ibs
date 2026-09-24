import os
import io
import json
import re
import requests
from datetime import datetime, date
from openpyxl import load_workbook


FILE_ID = os.environ["GOOGLE_DRIVE_FILE_ID"]
SERVICE_ACCOUNT_JSON = os.environ["GOOGLE_SERVICE_ACCOUNT_JSON"]

OUTPUT_FILE = "public/data.json"


def get_access_token(service_account_info):
    from cryptography.hazmat.primitives import hashes, serialization
    from cryptography.hazmat.primitives.asymmetric import padding
    import base64
    import time

    def b64url(data):
        return base64.urlsafe_b64encode(data).rstrip(b"=").decode()

    now = int(time.time())

    header = {
        "alg": "RS256",
        "typ": "JWT",
    }

    payload = {
        "iss": service_account_info["client_email"],
        "scope": "https://www.googleapis.com/auth/drive.readonly",
        "aud": "https://oauth2.googleapis.com/token",
        "iat": now,
        "exp": now + 3600,
    }

    header_b64 = b64url(
        json.dumps(header, separators=(",", ":")).encode()
    )

    payload_b64 = b64url(
        json.dumps(payload, separators=(",", ":")).encode()
    )

    unsigned = f"{header_b64}.{payload_b64}".encode()

    private_key = serialization.load_pem_private_key(
        service_account_info["private_key"].encode(),
        password=None,
    )

    signature = private_key.sign(
        unsigned,
        padding.PKCS1v15(),
        hashes.SHA256(),
    )

    jwt = f"{header_b64}.{payload_b64}.{b64url(signature)}"

    response = requests.post(
        "https://oauth2.googleapis.com/token",
        data={
            "grant_type": "urn:ietf:params:oauth:grant-type:jwt-bearer",
            "assertion": jwt,
        },
        timeout=60,
    )

    response.raise_for_status()
    return response.json()["access_token"]


def download_excel():
    service_account = json.loads(SERVICE_ACCOUNT_JSON)

    token = get_access_token(service_account)

    url = f"https://www.googleapis.com/drive/v3/files/{FILE_ID}"
    response = requests.get(
        url,
        headers={
            "Authorization": f"Bearer {token}"
        },
        params={
            "alt": "media"
        },
        timeout=120,
    )

    response.raise_for_status()

    return response.content


def normalize_header(value):
    if value is None:
        return ""

    return re.sub(
        r"\s+",
        " ",
        str(value).strip().lower()
    )


def parse_date(value):
    if value is None:
        return None

    if isinstance(value, datetime):
        return value.date()

    if isinstance(value, date):
        return value

    text = str(value).strip()

    if not text:
        return None

    # Format: dd/mm/yyyy
    formats = [
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y-%m-%d",
        "%d/%m/%y",
        "%d-%m-%y",
    ]

    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass

    # Indonesian month names
    months = {
        "januari": "01",
        "februari": "02",
        "maret": "03",
        "april": "04",
        "mei": "05",
        "juni": "06",
        "juli": "07",
        "agustus": "08",
        "september": "09",
        "oktober": "10",
        "november": "11",
        "desember": "12",
    }

    lowered = text.lower()

    for month_name, month_number in months.items():
        if month_name in lowered:
            match = re.search(
                r"(\d{1,2})\s+" + month_name + r"\s+(\d{4})",
                lowered,
            )

            if match:
                day = int(match.group(1))
                year = int(match.group(2))

                try:
                    return date(
                        year,
                        int(month_number),
                        day,
                    )
                except ValueError:
                    return None

    return None


def find_columns(ws):
    header_row = None
    name_col = None
    birthday_col = None
    generation_col = None

    # File format:
    # Row 2 = main headers
    # Row 3 = sub headers
    # Row 4+ = data
    #
    # We still search dynamically so small changes
    # in formatting don't break the script.

    for row in range(1, min(ws.max_row, 20) + 1):
        headers = [
            normalize_header(
                ws.cell(row=row, column=col).value
            )
            for col in range(1, ws.max_column + 1)
        ]

        for col, header in enumerate(headers, start=1):

            if header in ("nama siswa", "nama"):
                name_col = col

            elif "tanggal lahir" in header:
                birthday_col = col

            elif header == "angkatan":
                generation_col = col

        if name_col and birthday_col and generation_col:
            header_row = row
            break

    if not name_col:
        raise RuntimeError(
            "Kolom 'Nama Siswa' tidak ditemukan."
        )

    if not birthday_col:
        raise RuntimeError(
            "Kolom 'Tanggal Lahir' tidak ditemukan."
        )

    if not generation_col:
        raise RuntimeError(
            "Kolom 'ANGKATAN' tidak ditemukan."
        )

    return (
        header_row,
        name_col,
        birthday_col,
        generation_col,
    )


def load_students(excel_bytes):
    workbook = load_workbook(
        io.BytesIO(excel_bytes),
        data_only=True,
    )

    print(
        "Workbook sheets:",
        ", ".join(workbook.sheetnames)
    )

    # File baru menggunakan sheet MAIN
    if "MAIN" not in workbook.sheetnames:
        raise RuntimeError(
            "Sheet 'MAIN' tidak ditemukan. "
            f"Sheet yang tersedia: {', '.join(workbook.sheetnames)}"
        )

    ws = workbook["MAIN"]

    (
        header_row,
        name_col,
        birthday_col,
        generation_col,
    ) = find_columns(ws)

    print(f"Header row: {header_row}")
    print(f"Nama Siswa column: {name_col}")
    print(f"Tanggal Lahir column: {birthday_col}")
    print(f"ANGKATAN column: {generation_col}")

    students = []

    # Data dimulai setelah header/sub-header.
    # Dengan format sekarang, data efektif mulai row 4.
    for row in range(header_row + 2, ws.max_row + 1):

        name = ws.cell(
            row=row,
            column=name_col
        ).value

        birthday = ws.cell(
            row=row,
            column=birthday_col
        ).value

        generation = ws.cell(
            row=row,
            column=generation_col
        ).value

        if name is None:
            continue

        name = str(name).strip()

        if not name:
            continue

        birthday_date = parse_date(birthday)

        if not birthday_date:
            continue

        generation = (
            str(generation).strip()
            if generation is not None
            else ""
        )

        students.append({
            "nama": name,
            "tanggalLahir": birthday_date.isoformat(),
            "angkatan": generation,
        })

    return students


def main():
    print("Downloading Excel from Google Drive...")

    excel_bytes = download_excel()

    print("Excel downloaded successfully.")

    students = load_students(excel_bytes)

    os.makedirs(
        os.path.dirname(OUTPUT_FILE),
        exist_ok=True,
    )

   from datetime import datetime, timezone

    with open(
    OUTPUT_FILE,
    "w",
    encoding="utf-8",
) as file:
    json.dump(
        {
            "updatedAt": datetime.now(timezone.utc).isoformat(),
            "people": students,
        },
        file,
        ensure_ascii=False,
        indent=2,
    )

    print(
        f"Generated {OUTPUT_FILE}"
    )

    print(
        f"Total students: {len(students)}"
    )


if __name__ == "__main__":
    main()
